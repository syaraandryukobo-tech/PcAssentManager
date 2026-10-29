from PySide6.QtWidgets import QDialog,QMessageBox,QTableWidgetItem,QAbstractItemView
from PySide6.QtCore import Qt
from ui.frmPcInfo import Ui_Dialog
from datetime import datetime
import openpyxl
import os

class PcInfoForm(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.ui = Ui_Dialog()
        self.ui.setupUi(self)
        self.ui.lblManageNo.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.ui.txtStartDate.setAlignment(Qt.AlignRight | Qt.AlignVCenter)

        # コンボボックスの手入力可否
        self.ui.cmbLocation.setEditable(False)
        
        # 検索ボタン
        self.ui.btnSearch.clicked.connect(self.searchPcinfo)
        
        # 更新ボタン
        self.ui.btnUpdate.clicked.connect(self.updatePcInfo)
        
        # 閉じるボタン
        self.ui.btnClose.clicked.connect(self.closeWindow)
        
        # コンボボックス初期化
        self.initComboBox()
        
        # テーブル初期化
        self.initTable()
        
    def initComboBox(self):
        locations = [
            "",
            "東京本社",
            "大阪支社",
            "名古屋支社"
        ]
        
        # 拠点
        self.ui.cmbLocation.clear()
        self.ui.cmbLocation.addItems(locations)
        
    def searchPcinfo(self):
        location = self.ui.cmbLocation.currentText()
        userName =self.ui.txtUserName.text().strip()

        try:
            wb = openpyxl.load_workbook("Excel/pc_list.xlsx")
            ws = wb["パソコン発番表"]
            print("Excel読込成功")
            print("シート取得成功")

        except Exception as e:
            QMessageBox.critical(self,"エラー",f"Excel読込エラー\n{e}")

        print("拠点：", location)
        print("氏名：", userName)

        searchCount = 0

        # 検索結果クリア
        self.ui.tblPcInfo.setRowCount(0)

        for row in ws.iter_rows(min_row=2, values_only=True):
            # 拠点チェック
            if location != "" and row[6] != location:
                continue

            # 氏名チェック
            if userName != "" and userName not in str(row[4]):
                continue

            searchCount += 1

            print(row)

            rowPosition = self.ui.tblPcInfo.rowCount()
            self.ui.tblPcInfo.insertRow(rowPosition)

            manage_no = QTableWidgetItem(str(row[0]))
            manage_no.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self.ui.tblPcInfo.setItem(rowPosition, 0,manage_no)

            self.ui.tblPcInfo.setItem(rowPosition, 1,QTableWidgetItem(str(row[1])))
            self.ui.tblPcInfo.setItem(rowPosition, 2,QTableWidgetItem(str(row[4])))

            if row[12] is None:
                serial_no_text = ""
            else:
                serial_no_text = str(row[12])

            serial_no = QTableWidgetItem(serial_no_text)
            self.ui.tblPcInfo.setItem(rowPosition, 3,serial_no)

            if row[13] is None:
                start_date_text = ""
            elif hasattr(row[13], "strftime"):
                start_date_text = row[13].strftime("%Y/%m/%d")
            else:
                start_date_text = str(row[13])

            start_date = QTableWidgetItem(start_date_text)
            start_date.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self.ui.tblPcInfo.setItem(rowPosition, 4,start_date)

        totalCount =ws.max_row -1

        print("検索件数：", searchCount)
        print("総件数：", totalCount)

        self.ui.lblResultCount.setText(f"{searchCount}件／全{totalCount}件")

        wb.close()

    def initTable(self):
        headers = [
            "管理番号",
            "端末名",
            "氏名",
            "シリアル番号",
            "使用開始日"
        ]

        self.ui.tblPcInfo.setHorizontalHeaderLabels(headers)
        self.ui.tblPcInfo.setRowCount(0)
        self.ui.tblPcInfo.setColumnWidth(0, 60) #  管理番号
        self.ui.tblPcInfo.setColumnWidth(1, 110) # 端末名
        self.ui.tblPcInfo.setColumnWidth(2, 90) # 氏名
        self.ui.tblPcInfo.setColumnWidth(3, 100) # シリアル番号
        self.ui.tblPcInfo.setColumnWidth(4, 80) # 使用開始日

        self.ui.tblPcInfo.setSelectionBehavior(QAbstractItemView.SelectRows)

        self.ui.tblPcInfo.setEditTriggers(QAbstractItemView.NoEditTriggers)

        # 行選択時イベント
        self.ui.tblPcInfo.itemSelectionChanged.connect(self.onPcInfoSelected)

    def onPcInfoSelected(self):
        row = self.ui.tblPcInfo.currentRow()

        if row < 0:
            return

        manage_no = self.ui.tblPcInfo.item(row,0).text()
        self.selected_manage_no = manage_no

        print("管理番号：", manage_no)

        wb = openpyxl.load_workbook("Excel/pc_list.xlsx")
        ws = wb["パソコン発番表"]

        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=1).value) == manage_no:
                self.ui.lblManageNo.setText(manage_no)
                self.ui.lblTerminalName.setText(str(ws.cell(row=row, column=2).value))
                self.ui.lblUserName.setText(str(ws.cell(row=row, column=5).value))

                serial_no = ws.cell(row=row, column=13).value
                self.ui.txtSerialNo.setText("" if serial_no is None else str(serial_no))

                start_date = ws.cell(row=row, column=14).value
                if start_date is None:
                    start_date_text = ""
                elif hasattr(start_date, "strftime"):
                    start_date_text = start_date.strftime("%Y/%m/%d")
                else:
                    start_date_text = str(start_date)
                self.ui.txtStartDate.setText(start_date_text)

                remarks =ws.cell(row=row, column=15).value
                self.ui.txtRemarks.setPlainText("" if remarks is None else str(remarks))

    def updatePcInfo(self):
        print("更新対象")
        print(self.selected_manage_no)

        manage_no = self.selected_manage_no
        serial_no = self.ui.txtSerialNo.text()
        start_date = self.ui.txtStartDate.text()
        start_date_value = datetime.strptime(start_date, "%Y/%m/%d")
        remarks = self.ui.txtRemarks.toPlainText()

        print(manage_no)
        print(serial_no)
        print(start_date)
        print(remarks)

        excel_path = os.path.join("Excel","pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["パソコン発番表"]

        target_row = None

        for row in range(2, ws.max_row + 1):
            excel_manage_no = ws.cell(row=row,column=1).value

            print("Excel：", excel_manage_no, "選択：", manage_no)

            if str(excel_manage_no) == str(manage_no):
                target_row = row
                break

        if target_row is not None:
            ws.cell(row=target_row, column=13, value=serial_no)
            ws.cell(row=target_row, column=14, value=start_date_value)
            ws.cell(row=target_row, column=15, value=remarks)

            QMessageBox.information(self,"成功","更新できました。")

        wb.save(excel_path)    

        self.clearEditArea()
        self.searchPcinfo()

    def clearEditArea(self):
        self.ui.lblManageNo.setText("○○")
        self.ui.lblTerminalName.setText("○○○○○○○")
        self.ui.lblUserName.setText("○○○○○○")
        self.ui.txtSerialNo.clear()
        self.ui.txtStartDate.clear()
        self.ui.txtRemarks.clear()
        self.selected_manage_no = None

    def closeWindow(self):
        self.close()

        if self.parent():
            self.parent().show()


if __name__ == "__main__":
    import sys
    from PySide6.QtWidgets import QApplication

    app = QApplication(sys.argv)

    window = PcInfoForm()
    window.show()

    sys.exit(app.exec())
