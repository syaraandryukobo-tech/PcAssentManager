from PySide6.QtWidgets import QDialog ,QMessageBox
from PySide6.QtCore import QTimer
from ui.frmPcRegister import Ui_Dialog
import openpyxl
import os

class PcNewRegForm(QDialog):

    # 拠点コード辞書（定数扱い）
    location_map = {
        "東京本社": "HON",
        "大阪支社": "OSK",
        "名古屋支社": "NGY",
        "営業部": "HON" # 例外ルール
    }

    keitai_map = {
        "ノート": "N",
        "デスクトップ": "D"
    }

    def __init__(self, parent=None):
        super().__init__(parent)
        self.ui = Ui_Dialog()
        self.ui.setupUi(self)
        self.ui.lstUser.itemClicked.connect(self.onUserSelected)
        self.ui.lstPc.itemClicked.connect(self.onPcSelected)

        # 決定ボタン
        self.ui.btnDecision.clicked.connect(self.onRegisterClicked)
        #if self.selected_location_name == "営業部":
            #QMessageBox.information(self, "確認", "営業部は東京本社扱いとなります。")

        self.ui.btnClose.clicked.connect(self.close)
      
        # 行間を詰める設定
        self.ui.lstUser.setStyleSheet(
            "QListWidget {"
            "    background-color: white;"
            "    border: 1px solid #c0c0c0;"
            "}"
            "QListWidget::item {"
            "    padding: 2px 4px;"
            "    height: 18px;"
            "}"
        )

        # 拠点絞込
        self.ui.cmbUserLocation.clear()
        self.ui.cmbUserLocation.addItems([
            "東京本社",
            "大阪支社",
            "名古屋支社"
        ])

        # 拠点絞込が変更されたらユーザー一覧を更新
        self.ui.cmbUserLocation.currentTextChanged.connect(self.loadUsersByLocation)

        # 初期表示のために明示的に呼ぶ
        initial_location = self.ui.cmbUserLocation.currentText()
        self.loadUsersByLocation(initial_location)

        # 行間を詰める設定
        self.ui.lstPc.setStyleSheet(
            "QListWidget {"
            "    background-color: white;"
            "    border: 1px solid #c0c0c0;"
            "}"
            "QListWidget::item {"
            "    padding: 2px 4px;"
            "    height: 18px;"
            "}"
        )

        # 年度絞込（パソコンマスタのB列から重複なしで取得
        excel_path = os.path.join("Excel", "pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws =wb["パソコンマスタ"]

        years = set() # 重複排除"

        for row in range(2,ws.max_row + 1): # 1行目はヘッダー
            val = ws.cell(row=row, column=2).value # B列は年度"
            if isinstance(val, int): # 数値だけ採用（空欄や文字列は除外）
                years.add(val)

        # 降順にソート
        sorted_years = sorted(years, reverse=True)

        # コンボボックスへ設定
        self.ui.cmbPcYear.clear()
        self.ui.cmbPcYear.addItems([str(y) for y in sorted_years])

        # UIが描画された後に初期年度で読み込む
        QTimer.singleShot(0, lambda: self.loadPcsByYear(self.ui.cmbPcYear.currentText()))

        # 年度絞込が変更されたらパソコン一覧を更新
        self.ui.cmbPcYear.currentTextChanged.connect(self.loadPcsByYear)

        # 初期表示のために明示的に呼ぶ
        initial_location = self.ui.cmbPcYear.currentText()
        self.loadPcsByYear(initial_location)

    def loadUsersByLocation(self, location):

        # Excelファイルパス
        excel_path = os.path.join("Excel", "pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["ユーザーマスタ"]

        # リストを初期化
        self.ui.lstUser.clear()

        # 2行目から最終行まで走査"
        for row in range(2, ws.max_row + 1):
            user_location = ws.cell(row=row, column=4).value # 拠点(D列)
            user_code = ws.cell(row=row, column=1).value  # ユーザーコード(A列)
            user_name = ws.cell(row=row, column=2).value # 氏名(B列)

            # 拠点一致ならリストに追加
            if user_location == location:
                display_text = f"{user_code}-{user_name}"
                self.ui.lstUser.addItem(display_text) 

    def loadPcsByYear(self, year):

        # コンボボックスの値は strなので intに変換
        try:
            year = int(year)
        except ValueError:
            return

        # Excelファイルパス
        excel_path = os.path.join("Excel", "pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["パソコンマスタ"]

        # リストを初期化
        self.ui.lstPc.clear()

        # 2行目から最終行まで走査"
        for row in range(2, ws.max_row + 1):
            pc_year = ws.cell(row=row, column=2).value # 年度(B列)
            pc_keitai = ws.cell(row=row, column=6).value  # 形態(F列)
            pc_maker = ws.cell(row=row, column=3).value # メーカー(C列)
            pc_kisyu = ws.cell(row=row, column=4).value # 機種(D列)

            # 年度一致ならリストに追加
            if pc_year == year:
                display_text = f"{pc_keitai}-{pc_maker}-{pc_kisyu}"
                self.ui.lstPc.addItem(display_text)

        self.selected_year = int(year) # ここで年度を保持
        print("年度:", year)

    def getLocationCode(self, location_name, department):
        # 営業部は例外ルール
        if department == "営業部":
            return "HON"
        
        return self.location_map.get(location_name, "")

    def onUserSelected(self, item):
        # item.text()は「S-22001-ユーザー1」
        parts = item.text().split("-")
        user_code = parts[0] + "-" + parts[1]        

        # Excelから拠点を取得
        excel_path = os.path.join("Excel", "pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["ユーザーマスタ"]

        for row in range(2, ws.max_row + 1):
            code = ws.cell(row=row, column=1).value
            if code == user_code:
                user_name = ws.cell(row=row, column=2).value # 氏名(B列)
                department = ws.cell(row=row, column=3).value # 所属(C列)
                location_name = ws.cell(row=row, column=4).value # 拠点(D列)
                mail_address = ws.cell(row=row, column=6).value # メールアドレス(F列)
                
                self.selected_user_code = code
                self.selected_user_name = user_name
                self.selected_department = department
                self.selected_location_name = location_name
                self.selected_user_mail = mail_address

                print("ユーザーコード:", code)
                print("氏名:", user_name)
                print("所属:", department)
                print("拠点名:", location_name)
                print("メールアドレス:", mail_address)
                break

    def getkeitaiCode(self, keitai_name):
        return self.keitai_map.get(keitai_name, "")

    def onPcSelected(self, item):
        parts = item.text().split("-")
        keitai_name = parts[0] # ノート or デスクトップ
        maker = parts[1]
        model = parts[2]

        print("形態",keitai_name)
        print("メーカー", maker)
        print("機種", model)

        excel_path = os.path.join("Excel", "pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["パソコンマスタ"]

        for row in range(2, ws.max_row + 1):
            excel_keitai = ws.cell(row=row, column=6).value
            excel_maker = ws.cell(row=row, column=3).value
            excel_model =ws.cell(row=row, column=4).value
            excel_year = ws.cell(row=row, column=2).value

            if (excel_keitai == keitai_name and
                excel_maker == maker and
                excel_model == model and
                excel_year == self.selected_year):
                pc_code = ws.cell(row=row, column=1).value
                type_no = ws.cell(row=row, column=5).value

                keitai_code = self.getkeitaiCode(keitai_name)              

                self.selected_pd_code =pc_code
                self.selected_maker = maker
                self.selected_model = model
                self.selected_type_no = type_no
                self.selected_keitai_name = keitai_name
                self.selected_keitai_code = keitai_code
            
                print("パソコンコード:", pc_code)
                print("メーカー:", maker)
                print("機種:", model)
                print("型番:", type_no)
                print("形態名", keitai_name)
                print("形態コード", keitai_code)
                break

    def getNextNumber(self, location_code,keitai_code):
        excl_path = os.path.join("Excel", "pc_list.xlsx")
        wb = openpyxl.load_workbook(excl_path)
        ws = wb["パソコン発番表"]

        prefix = f"ABC-{location_code}-{keitai_code}"

        numbers = []

        for row in range(2, ws.max_row + 1):
            pc_name = ws.cell(row=row, column=2).value # 端末名が入っている列

            if pc_name and pc_name.startswith(prefix):
                num_str = pc_name.replace(prefix, "")
                if num_str.isdigit():
                    numbers.append(int(num_str))

        next_num = max(numbers) + 1 if numbers else 1

        return f"{next_num:03d}"

    def onRegisterClicked(self):

        # 営業部の確認イメージ
        if self.selected_department == "営業部":
            QMessageBox.information(self, "確認","営業部は東京本社扱いとなります。")

        # 拠点コードと形態コード
        location_code = self.getLocationCode(
            self.selected_location_name,
            self.selected_department
        )
        keitai_code = self.selected_keitai_code

        # 連番採番
        next_number = self.getNextNumber(location_code, keitai_code)

        # 端末名生成
        terminal_name = f"ABC-{location_code}-{keitai_code}{next_number}"

        print("生成された端末名:", terminal_name)

        # Excel登録処置
        self.registerToExcel(terminal_name)

        # 登録完了メッセージ
        QMessageBox.information(self, "成功", "登録できました。")

        # フォームを閉じる
        self.close()

    def getNextKanriNumber(self):
        excel_path = os.path.join("Excel", "pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["パソコン発番表"]

        numbers = []

        for row in range(2, ws.max_row + 1):
            kanri = ws.cell(row=row, column=1).value # 管理番号(A列)
            if kanri and str(kanri).isdigit():
                numbers.append(int(kanri))

        next_num = max(numbers) + 1 if numbers else 1
        return next_num

    def registerToExcel(self, terminal_name):
        excel_path = os.path.join("Excel", "pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["パソコン発番表"]

        # 管理番号採番
        kanri_number = self.getNextKanriNumber()

        # 次の空行
        next_row = ws.max_row + 1

        # 書き込み
        ws.cell(row=next_row, column=1).value = kanri_number
        ws.cell(row=next_row, column=2).value = terminal_name
        ws.cell(row=next_row, column=3).value = self.selected_user_code
        ws.cell(row=next_row, column=4).value = self.selected_user_mail
        ws.cell(row=next_row, column=5).value = self.selected_user_name
        ws.cell(row=next_row, column=6).value = self.selected_department
        ws.cell(row=next_row, column=7).value = self.selected_location_name
        ws.cell(row=next_row, column=8).value = self.selected_pd_code
        ws.cell(row=next_row, column=9).value = self.selected_maker
        ws.cell(row=next_row, column=10).value = self.selected_model
        ws.cell(row=next_row, column=11).value = self.selected_type_no
        ws.cell(row=next_row, column=12).value = self.selected_keitai_name

        wb.save(excel_path)

 
    def closeEvent(self, event):
        if self.parent():
            self.parent().showNormal()
        event.accept()

