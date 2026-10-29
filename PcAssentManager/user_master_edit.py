from PySide6.QtWidgets import QDialog,QMessageBox,QTableWidgetItem,QAbstractItemView
from ui.frmUserMasterEdit import Ui_Dialog
import openpyxl
import os

class UserMasterEditForm(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.ui = Ui_Dialog()
        self.ui.setupUi(self)      

        # コンボボックスの手入力可否
        self.ui.cmbLocation.setEditable(False)

        # 検索ボタン
        self.ui.btnSearch.clicked.connect(self.searchUser)

        # 更新ボタン
        self.ui.btnUpdate.clicked.connect(self.updateUser)

        # 閉じるボタン
        self.ui.btnClose.clicked.connect(self.closeWindow)

        # コンボボックス初期化
        self.initComboBox()

        # テーブル初期化
        self.initTable()

    def initComboBox(self):
        locations = [
            "",
            "東京本社",
            "大阪支社",
            "名古屋支社"
        ]

        departments = [
            "",
            "営業部",
            "総務部",
            "経理部",
            "人事部",
            "システム部"
        ]

        # 拠点
        self.ui.cmbLocation.clear()
        self.ui.cmbLocation.addItems(locations)

        self.ui.cmbEditLocation.clear()
        self.ui.cmbEditLocation.addItems(locations)

        # 所属
        self.ui.cmbEditDepartment.clear()
        self.ui.cmbEditDepartment.addItems(departments)

    def searchUser(self):
        location = self.ui.cmbLocation.currentText()
        userName =self.ui.txtUserName.text().strip()

        if location == "" and userName == "":
            QMessageBox.information(self,"警告","検索条件を入力してください。")
            return

        try:
            wb = openpyxl.load_workbook("Excel/pc_list.xlsx")
            ws = wb["ユーザーマスタ"]
            print("Excel読込成功")
            print("シート取得成功")

        except Exception as e:
            QMessageBox.critical(self,"エラー",f"Excel読込エラー\n{e}")

        print("拠点：", location)
        print("氏名：", userName)

        searchCount = 0

        # 検索結果クリア
        self.ui.tblUserMaster.setRowCount(0)

        for row in ws.iter_rows(min_row=2, values_only=True):
            # 拠点チェック
            if location != "" and row[3] != location:
                continue

            # 氏名チェック
            if userName != "" and userName not in str(row[1]):
                continue

            searchCount += 1

            print(row)

            rowPosition = self.ui.tblUserMaster.rowCount()
            self.ui.tblUserMaster.insertRow(rowPosition)

            self.ui.tblUserMaster.setItem(rowPosition, 0,QTableWidgetItem(str(row[0])))
            self.ui.tblUserMaster.setItem(rowPosition, 1,QTableWidgetItem(str(row[3])))
            self.ui.tblUserMaster.setItem(rowPosition, 2,QTableWidgetItem(str(row[2])))
            self.ui.tblUserMaster.setItem(rowPosition, 3,QTableWidgetItem(str(row[1])))
            self.ui.tblUserMaster.setItem(rowPosition, 4,QTableWidgetItem(str(row[5])))

        totalCount =ws.max_row -1

        print("検索件数：", searchCount)
        print("総件数：", totalCount)

        self.ui.lblResultCount.setText(f"{searchCount}件／全{totalCount}件")

        wb.close()

    def initTable(self):
        headers = [
            "ユーザーコード",
            "拠点",
            "所属",
            "氏名",
            "メールアドレス"
        ]

        self.ui.tblUserMaster.setHorizontalHeaderLabels(headers)
        self.ui.tblUserMaster.setRowCount(0)
        self.ui.tblUserMaster.setColumnWidth(0, 80) # ユーザーコード
        self.ui.tblUserMaster.setColumnWidth(1, 90) # 拠点
        self.ui.tblUserMaster.setColumnWidth(2, 80) # 所属
        self.ui.tblUserMaster.setColumnWidth(3, 100) # 氏名
        self.ui.tblUserMaster.setColumnWidth(4, 140) # メールアドレス

        self.ui.tblUserMaster.setSelectionBehavior(QAbstractItemView.SelectRows)

        self.ui.tblUserMaster.setEditTriggers(QAbstractItemView.NoEditTriggers)

        # 行選択時イベント
        self.ui.tblUserMaster.itemSelectionChanged.connect(self.onUserSelected)

    def onUserSelected(self):
        row = self.ui.tblUserMaster.currentRow()

        if row < 0:
            return

        user_code = self.ui.tblUserMaster.item(row, 0).text()
        location = self.ui.tblUserMaster.item(row, 1).text()
        department = self.ui.tblUserMaster.item(row, 2).text()
        user_name = self.ui.tblUserMaster.item(row, 3).text()
        mail_address = self.ui.tblUserMaster.item(row, 4).text()

        print("ユーザーコード：", user_code)
        print("拠点：", location)
        print("所属：", department)
        print("氏名：", user_name)
        print("メールアドレス：", mail_address) 

        self.selected_user_code =user_code
        self.ui.cmbEditLocation.setCurrentText(location)
        self.ui.cmbEditDepartment.setCurrentText(department)
        self.ui.txtEditUserName.setText(user_name)
        self.ui.txtMailAddress.setText(mail_address)

    def updateUser(self):
        print("更新対象")
        print(self.selected_user_code)

        user_code = self.selected_user_code
        location = self.ui.cmbEditLocation.currentText()
        department =self.ui.cmbEditDepartment.currentText()
        user_name = self.ui.txtEditUserName.text()
        mail_address = self.ui.txtMailAddress.text()

        print(location)
        print(department)
        print(user_name)
        print(mail_address)

        excel_path = os.path.join("Excel","pc_list.xlsx")
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["ユーザーマスタ"]

        target_row = None

        for row in range(2, ws.max_row + 1):
            excel_user_code = ws.cell(row=row,column=1).value

            if excel_user_code == user_code:
                target_row = row
                break

        if target_row is not None:
            ws.cell(row=target_row, column=2, value=user_name)
            ws.cell(row=target_row, column=3, value=department)
            ws.cell(row=target_row, column=4, value=location)
            ws.cell(row=target_row, column=6, value=mail_address)

            QMessageBox.information(self,"成功","更新できました。")

        pc_ws = wb["パソコン発番表"]

        for row in range(2, pc_ws.max_row + 1):
            pc_user_code = pc_ws.cell(row=row, column=3).value

            if pc_user_code == user_code:
                pc_ws.cell(row=row, column=4,value=mail_address)
                pc_ws.cell(row=row, column=5, value=user_name)
                pc_ws.cell(row=row, column=6, value=department)
                pc_ws.cell(row=row, column=7, value=location)

        wb.save(excel_path)    

        self.clearEditArea()
        self.searchUser()

    def clearEditArea(self):
        print("clearEditArea 実行")

        self.ui.cmbEditLocation.setCurrentIndex(0)
        self.ui.cmbEditDepartment.setCurrentIndex(0)
        self.ui.txtEditUserName.clear()
        self.ui.txtMailAddress.clear()
        self.selected_user_code = None

        print("クリア後氏名：", self.ui.txtEditUserName.text())
    
    def closeWindow(self):
        self.close()

        if self.parent():
            self.parent().show()

if __name__ == "__main__":
    import sys
    from PySide6.QtWidgets import QApplication

    app = QApplication(sys.argv)

    window = UserMasterEditForm()
    window.show()

    sys.exit(app.exec())