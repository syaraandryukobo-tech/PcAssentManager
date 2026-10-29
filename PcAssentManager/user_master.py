from PySide6.QtWidgets import QDialog, QMessageBox
from PySide6.QtCore import Qt, QEvent
from ui.frmUserMasterRegister import Ui_Dialog
import openpyxl
import os
    
class UserMasterForm(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.ui = Ui_Dialog()
        self.ui.setupUi(self)
        self.ui.txtJoinYear.editingFinished.connect(self.generateUserCode)
        self.ui.txtUserCode.setFocusPolicy(Qt.NoFocus)
        self.ui.txtUserCode.setReadOnly(True)
        self.ui.txtUserName.installEventFilter(self)
        self.ui.txtMailAddress.setFocusPolicy(Qt.ClickFocus)
        self.ui.txtMailAddress.mousePressEvent = self.emailClicked     
     
        # コンボボックスの手入力可否
        self.ui.cmbDepartment.setEditable(False)
        self.ui.cmbLocation.setEditable(False)
        self.ui.cmbEmploymentType.setEditable(False)

        # 登録ボタンのイベント
        self.ui.btnRegister.clicked.connect(self.registerUser)

        # 閉じるボタン
        self.ui.btnClose.clicked.connect(self.close)

        # 雇用区分
        self.ui.cmbEmploymentType.clear()
        self.ui.cmbEmploymentType.addItem("")
        self.ui.cmbEmploymentType.addItems([
            "正社員",
            "契約社員",
            "派遣"
        ])

        # 拠点
        self.ui.cmbLocation.clear()
        self.ui.cmbLocation.addItem("")
        self.ui.cmbLocation.addItems([
            "東京本社",
            "大阪支社",
            "名古屋支社"
        ])

        # 所属
        self.ui.cmbDepartment.clear()
        self.ui.cmbDepartment.addItem("")
        self.ui.cmbDepartment.addItems([
            "営業部",
            "総務部",
            "経理部",
            "人事部",
            "システム部"
        ])
        
    def keyPressEvent(self, event):
        if event.key() in (Qt.Key_Return, Qt.Key_Enter):
            event.ignore()
            return
        super().keyPressEvent(event)

    def generateUserCode(self):

        # 雇用区分からプレフィックスを決定
        empType = self.ui.cmbEmploymentType.currentText()

        if empType in ["正社員", "契約社員"]:
            prefix = "S-"
        elif empType == "派遣":
            prefix ="H-"
        else:
            QMessageBox.warning(self, "エラー", "雇用区分が不正です。")

        # 入社年度を取得（下2桁を使う）
        join_year = self.ui.txtJoinYear.text().strip()

        if not join_year.isdigit() or len(join_year) !=4:
            QMessageBox.warning(self, "エラー", "入社年度は4桁の数字で入力してください。")
            return

        year2 = join_year[-2:] # 下2桁を使う

        # Excelファイルのパス
        excel_path = os.path.join("Excel", "pc_list.xlsx")

        # Excelを開く
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["ユーザーマスタ"]

        # A列（ユーザーコード）を走査して最大番号を探す
        max_number = 0

        for row in range(2, ws.max_row + 1): # 1行目はヘッダー
            code = ws.cell(row=row, column=1).value

            if code and code.startswith(prefix + year2):
                number = int(code[-3:]) # 末尾3桁を取得
                if number > max_number:
                    max_number = number

        # 新しい番号を生成
        new_number = max_number + 1
        new_code = f"{prefix}{year2}{new_number:03d}"

        # フォームにセット
        self.ui.txtUserCode.setText(new_code)

    def eventFilter(self, obj, event):
         # 氏名欄でのキー操作を監視
         if obj is self.ui.txtUserName and event.type() == QEvent.KeyPress:

             # Tabキー無効化（メールアドレス欄に移動させない）
             if event.key() == Qt.Key_Tab:
                 return True # ←フォーカス移動を止める

             # Shift + Tab(Backtab)無効化
             # if event.key() == Qt.Key_Backtab:
                 # return True

         # それ以外は通常処理
         return super().eventFilter(obj, event)     

    def emailClicked(self, event):
        current = self.ui.txtMailAddress.text().strip()

        # まだ何も入力されていない場合だけ自動補填
        if current == "":
            domain ="@example.co.jp"
            self.ui.txtMailAddress.setText(domain)

            # カーソルを「@」の前に移動させる
            at_index = domain.index("@")
            self.ui.txtMailAddress.setCursorPosition(at_index)

        # 元のイベントを呼び出す
        return super(type(self.ui.txtMailAddress),self.ui.txtMailAddress).mousePressEvent(event)
    
    def is_zenkaku(self, text):
            import unicodedata
            return any(unicodedata.east_asian_width(ch) in ['W', 'F'] for ch in text)
    
    def registerUser(self):
            year = self.ui.txtJoinYear.text().strip()
            mail = self.ui.txtMailAddress.text().strip()

            error_found = False
    
            if self.is_zenkaku(year):
                QMessageBox.warning(self, "入力エラー", "入社年度は半角数字4桁で入力してください。")
                error_found = True
    
            if self.is_zenkaku(mail):
                QMessageBox.warning(self, "入力エラー", "メールアドレスは半角で入力してください。")
                error_found = True

            # 「@example.co.jp」だけの場合エラー
            if mail == "@example.co.jp":
                QMessageBox.warning(self, "入力エラー", "@より前にメール名を入力してください。")
                error_found = True

            # どこかエラーがあれば登録処理は中断
            if error_found:
                return
            
            # ここまで来たら登録可能
            # QMessageBox.information(self, "確認", "入力チェックOK。登録処理を続行できます。")
      
            # 入力値を取得
            code = self.ui.txtUserCode.text()
            name = self.ui.txtUserName.text()
            dept = self.ui.cmbDepartment.currentText()
            base = self.ui.cmbLocation.currentText()
            empType = self.ui.cmbEmploymentType.currentText()
            year = self.ui.txtJoinYear.text()
            email = self.ui.txtMailAddress.text()

            # printで確認
            print("===== debug start =====")
            print("ユーザーコード", code)
            print("氏名:", name)
            print("所属:", dept)
            print("拠点:", base)
            print("雇用区分:", empType)
            print("入社年度:", year)
            print("メールアドレス:", email)
            print("===== debug end =====")

            # Excelファイルのパス
            excel_path = os.path.join("Excel", "pc_list.xlsx")

            # Excelを開く
            wb = openpyxl.load_workbook(excel_path)
            ws = wb["ユーザーマスタ"] # ←シート名

            # ユーザーコードが空の行を探す
            write_row = None
            for row in range(2,ws.max_row + 1): #1行目はヘッダーなので2行目から
                if ws.cell(row=row,column=1).value in (None, ""):
                    write_row = row
                    break

            # 空行が見つからなければ一番下に追加
            if write_row is None:
                write_row = ws.max_row + 1

            # Excelに書き込む（入社年度は不要）
            ws.cell(row=write_row, column=1, value=code)
            ws.cell(row=write_row, column=2, value=name)
            ws.cell(row=write_row, column=3, value=dept)
            ws.cell(row=write_row, column=4, value=base)
            ws.cell(row=write_row, column=5, value=empType)
            ws.cell(row=write_row, column=6, value=email)

            # 保存
            wb.save(excel_path)

            QMessageBox.information(self, "成功","登録できました。")

            # Excel保存後にフォームをクリア
            self.ui.txtUserCode.clear()
            self.ui.txtUserName.clear()
            self.ui.txtJoinYear.clear()
            self.ui.cmbDepartment.setCurrentIndex(0)
            self.ui.cmbLocation.setCurrentIndex(0)
            self.ui.cmbEmploymentType.setCurrentIndex(0)
            self.ui.txtMailAddress.clear()

            # カーソルを最初の項目に戻す
            self.ui.cmbEmploymentType.setFocus()

    def closeEvent(self, event):
        if self.parent():
            self.parent().showNormal()
        event.accept()
    
