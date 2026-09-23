from PySide6.QtWidgets import QDialog, QMessageBox
from PySide6.QtCore import Qt
from ui.frmPcMasterRegister import Ui_Dialog
import openpyxl
import os

class PcMasterForm(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.ui = Ui_Dialog()
        self.ui.setupUi(self)
        self.ui.txtYear.editingFinished.connect(self.generatePcCode)
        self.ui.txtPcCode.setFocusPolicy(Qt.NoFocus)
        self.ui.txtPcCode.setReadOnly(True)

        # 登録ボタン
        self.ui.btnRegister.clicked.connect(self.registerPc)

        # 閉じるボタン
        self.ui.btnClose.clicked.connect(self.close)

        # メーカー
        self.ui.cmbMaker.clear()
        self.ui.cmbMaker.addItem("")
        self.ui.cmbMaker.addItems([
            "Lenovo",
            "Dynabook",
            "HP"
            ])

        # 形態
        self.ui.cmbPcType.clear()
        self.ui.cmbPcType.addItem("")
        self.ui.cmbPcType.addItems([
            "ノート",
            "デスクトップ"
        ])

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key_Return, Qt.Key_Enter):
            event.ignore()
            return
        super().keyPressEvent(event)

    def generatePcCode(self):
    
        # 年度を取得（下2桁を使う）
        year = self.ui.txtYear.text().strip()
    
        if not year.isdigit() or len(year) !=4:
            QMessageBox.warning(self, "エラー", "年度は4桁の数字で入力してください。")
            return
    
        year2 = year[-2:] # 下2桁を使う
    
        # Excelファイルのパス
        excel_path = os.path.join("Excel", "pc_list.xlsx")
    
        # Excelを開く
        wb = openpyxl.load_workbook(excel_path)
        ws = wb["パソコンマスタ"]
    
        # A列（パソコンコード）を走査して最大番号を探す
        max_number = 0
    
        for row in range(2, ws.max_row + 1): # 1行目はヘッダー
            code = ws.cell(row=row, column=1).value
    
            if code and code.startswith("PC" + year2):
                number = int(code[-2:]) # 末尾2桁を取得
                if number > max_number:
                     max_number = number
    
        # 新しい番号を生成
        new_number = max_number + 1
        new_code = f"PC{year2}{new_number:02d}"
    
        # フォームにセット
        self.ui.txtPcCode.setText(new_code)

    def is_zenkaku(self, text):
        import unicodedata
        return any(unicodedata.east_asian_width(ch) in ['W', 'F'] for ch in text)
    
    def registerPc(self):
            year = self.ui.txtYear.text().strip()
            kisyu = self.ui.txtModel.text().strip()
            kataban = self.ui.txtModelNo.text().strip()

            error_found = False
    
            if self.is_zenkaku(year):
                QMessageBox.warning(self, "入力エラー", "年度は半角数字4桁で入力してください。")
                error_found = True
    
            if self.is_zenkaku(kisyu):
                QMessageBox.warning(self, "入力エラー", "機種は半角で入力してください。")
                error_found = True

            if self.is_zenkaku(kataban):
                QMessageBox.warning(self, "入力エラー", "型番は半角で入力してください。")
                error_found = True

            # どこかエラーがあれば登録処理は中断
            if error_found:
                return
            
            # ここまで来たら登録可能
            # QMessageBox.information(self, "確認", "入力チェックOK。登録処理を続行できます。")

            # 入力値を取得
            code = self.ui.txtPcCode.text()
            year = int(self.ui.txtYear.text())
            maker = self.ui.cmbMaker.currentText()
            kisyu = self.ui.txtModel.text()
            kataban = self.ui.txtModelNo.text()
            keitai = self.ui.cmbPcType.currentText()
            
            # printで確認
            print("===== debug start =====")
            print("パソコンコード", code)
            print("年度:", year)
            print("メーカー:", maker)
            print("機種:", kisyu)
            print("型番:", kataban)
            print("携帯:", keitai)
            print("===== debug end =====")
            
            # Excelファイルのパス
            excel_path = os.path.join("Excel", "pc_list.xlsx")
            
            # Excelを開く
            wb = openpyxl.load_workbook(excel_path)
            ws = wb["パソコンマスタ"] # ←シート名
            
            # パソコンコードが空の行を探す
            write_row = None
            for row in range(2,ws.max_row + 1): #1行目はヘッダーなので2行目から
                if ws.cell(row=row,column=1).value in (None, ""):
                    write_row = row
                    break
            
            # 空行が見つからなければ一番下に追加
            if write_row is None:
                write_row = ws.max_row + 1
            
            # Excelに書き込む
            ws.cell(row=write_row, column=1, value=code)
            ws.cell(row=write_row, column=2, value=year)
            ws.cell(row=write_row, column=3, value=maker)
            ws.cell(row=write_row, column=4, value=kisyu)
            ws.cell(row=write_row, column=5, value=kataban)
            ws.cell(row=write_row, column=6, value=keitai)
            
            # 保存
            wb.save(excel_path)
            
            QMessageBox.information(self, "成功","登録できました。")
            
            # Excel保存後にフォームをクリア
            self.ui.txtPcCode.clear()
            self.ui.txtYear.clear()
            self.ui.cmbMaker.setCurrentIndex(0)
            self.ui.txtModel.clear()
            self.ui.txtModelNo.clear()
            self.ui.cmbPcType.setCurrentIndex(0)         
            
            # カーソルを最初の項目に戻す
            self.ui.txtYear.setFocus()

    def closeEvent(self, event):
        if self.parent():
            self.parent().showNormal()
        event.accept()
